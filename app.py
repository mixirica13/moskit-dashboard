import os
import io
import sys
import json
import datetime
import requests
from flask import Flask, render_template, request, send_file, jsonify, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Carregar .env manualmente (sem dependência extra)
env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")
if os.path.exists(env_path):
    with open(env_path) as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                key, val = line.split("=", 1)
                os.environ.setdefault(key.strip(), val.strip())

app = Flask(__name__)
app.config["TEMPLATES_AUTO_RELOAD"] = True

API_BASE = "https://api.ms.prod.moskit.services/v2"
DEFAULT_API_KEY = os.environ.get("MOSKIT_API_KEY", "")


def moskit_headers(api_key):
    return {
        "apikey": api_key,
        "Content-Type": "application/json",
    }


def buscar_campos_pesquisa(api_key):
    """Busca os campos disponíveis para pesquisa de deals."""
    resp = requests.get(
        f"{API_BASE}/deals/search",
        headers=moskit_headers(api_key),
        timeout=30,
    )
    resp.raise_for_status()
    return resp.json()


def buscar_custom_fields(api_key):
    """Busca todos os custom fields e retorna um mapa id -> nome."""
    all_fields = []
    next_page_token = None

    while True:
        params = {"quantity": 200}
        if next_page_token:
            params["nextPageToken"] = next_page_token
        else:
            params["start"] = 0

        resp = requests.get(
            f"{API_BASE}/customFields",
            headers=moskit_headers(api_key),
            params=params,
            timeout=30,
        )
        resp.raise_for_status()

        fields = resp.json()
        if not fields:
            break

        all_fields.extend(fields)

        next_page_token = resp.headers.get("X-Moskit-Listing-Next-Page-Token")
        if not next_page_token:
            break

    return {f["id"]: f["name"] for f in all_fields}


def buscar_stages(api_key):
    """Busca todas as stages (etapas) e retorna um mapa id -> nome."""
    all_stages = []
    next_page_token = None

    while True:
        params = {"quantity": 200}
        if next_page_token:
            params["nextPageToken"] = next_page_token
        else:
            params["start"] = 0

        resp = requests.get(
            f"{API_BASE}/stages",
            headers=moskit_headers(api_key),
            params=params,
            timeout=30,
        )
        resp.raise_for_status()

        stages = resp.json()
        if not stages:
            break

        all_stages.extend(stages)

        next_page_token = resp.headers.get("X-Moskit-Listing-Next-Page-Token")
        if not next_page_token:
            break

    return {s["id"]: s.get("name", str(s["id"])) for s in all_stages}


def buscar_deals(api_key, conditions, quantity=200):
    """Busca deals usando o endpoint de pesquisa com paginação."""
    all_deals = []
    next_page_token = None

    while True:
        params = {"quantity": quantity}
        if next_page_token:
            params["nextPageToken"] = next_page_token
        else:
            params["start"] = 0

        resp = requests.post(
            f"{API_BASE}/deals/search",
            headers=moskit_headers(api_key),
            json=conditions,
            params=params,
            timeout=60,
        )
        resp.raise_for_status()

        deals = resp.json()
        if not deals:
            break

        all_deals.extend(deals)

        next_page_token = resp.headers.get("X-Moskit-Listing-Next-Page-Token")
        if not next_page_token:
            break

    return all_deals


def gerar_excel(deals, cf_map, titulo="Leads Moskit"):
    """Gera um arquivo Excel com os deals incluindo custom fields."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]  # max 31 chars for sheet name

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    cf_header_fill = PatternFill(start_color="8E44AD", end_color="8E44AD", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Colunas fixas
    colunas_fixas = [
        ("ID", "id"),
        ("Nome", "name"),
        ("Status", "status"),
        ("Motivo de Perda", "_lost_reason"),
        ("Fase", "stage.id"),
        ("Responsável", "responsible.id"),
        ("Contato", "_contact_id"),
        ("Data de Criação", "dateCreated"),
        ("Origem", "origin"),
        ("Source", "source"),
    ]

    # Descobrir todos os custom fields presentes nos deals
    cf_ids_presentes = []
    cf_ids_vistos = set()
    for deal in deals:
        for ecf in deal.get("entityCustomFields") or []:
            cf_id = ecf.get("id")
            if cf_id and cf_id not in cf_ids_vistos:
                cf_ids_vistos.add(cf_id)
                cf_ids_presentes.append(cf_id)

    total_colunas = len(colunas_fixas) + len(cf_ids_presentes)

    # Escrever cabeçalhos fixos
    for col_idx, (titulo, _) in enumerate(colunas_fixas, 1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Escrever cabeçalhos de custom fields
    for i, cf_id in enumerate(cf_ids_presentes):
        col_idx = len(colunas_fixas) + i + 1
        cf_nome = cf_map.get(cf_id, cf_id)
        cell = ws.cell(row=1, column=col_idx, value=cf_nome)
        cell.font = header_font
        cell.fill = cf_header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Escrever dados
    for row_idx, deal in enumerate(deals, 2):
        # Colunas fixas
        for col_idx, (_, campo) in enumerate(colunas_fixas, 1):
            valor = extrair_valor(deal, campo)
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

        # Custom fields
        cf_valores = {}
        for ecf in deal.get("entityCustomFields") or []:
            cf_valores[ecf.get("id")] = ecf.get("textValue", "")

        for i, cf_id in enumerate(cf_ids_presentes):
            col_idx = len(colunas_fixas) + i + 1
            cell = ws.cell(row=row_idx, column=col_idx, value=cf_valores.get(cf_id, ""))
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Ajustar largura das colunas
    for col_idx in range(1, total_colunas + 1):
        max_length = 0
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws.max_row):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def gerar_json_leads(deals, cf_map):
    """Gera JSON no formato [{meta_leadgen_id, nome, moskit_id}] para uma lista de deals."""
    # Inverter cf_map: nome -> id
    nome_para_id = {v: k for k, v in cf_map.items()}
    cf_leadgen_id = nome_para_id.get("meta_leadgen_id")

    resultado = []
    for deal in deals:
        # Buscar meta_leadgen_id nos custom fields do deal
        meta_leadgen = ""
        if cf_leadgen_id:
            for ecf in deal.get("entityCustomFields") or []:
                if ecf.get("id") == cf_leadgen_id:
                    meta_leadgen = ecf.get("textValue", "") or ""
                    break

        resultado.append({
            "meta_leadgen_id": meta_leadgen,
            "nome": deal.get("name", ""),
            "moskit_id": str(deal.get("id", "")),
        })

    return resultado


def extrair_valor(deal, campo):
    """Extrai valor de um deal, suportando campos aninhados e campos especiais."""
    if campo == "_contact_id":
        contacts = deal.get("contacts") or []
        if contacts:
            return contacts[0].get("id", "")
        return ""

    if campo == "_lost_reason":
        lost_reason = deal.get("lostReason")
        if isinstance(lost_reason, dict):
            return lost_reason.get("name", "")
        return lost_reason or ""

    partes = campo.split(".")
    valor = deal
    for parte in partes:
        if isinstance(valor, dict):
            valor = valor.get(parte)
        else:
            return ""
    return valor if valor is not None else ""


@app.route("/")
def index():
    return render_template("index.html", default_api_key=DEFAULT_API_KEY)


@app.route("/campos", methods=["POST"])
def campos():
    """Retorna os campos disponíveis para pesquisa."""
    data = request.json
    api_key = data.get("api_key", "").strip()
    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    try:
        campos = buscar_campos_pesquisa(api_key)
        return jsonify(campos)
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def montar_conditions(data):
    """Monta as condições de pesquisa a partir dos parâmetros.
    Retorna (conditions, filtros_locais) onde filtros_locais são filtros
    que precisam ser aplicados após a busca (ex: not_contains).
    """
    conditions = []
    filtros_locais = []  # filtros aplicados após busca (not_contains)
    campo_filtro = data.get("campo_filtro", "dateCreated")
    data_inicio = data.get("data_inicio", "").strip()
    data_fim = data.get("data_fim", "").strip()

    if data_inicio and data_fim:
        # gt é exclusivo, então usamos 1 segundo antes da meia-noite do dia anterior
        dt_inicio = datetime.datetime.strptime(data_inicio, "%Y-%m-%d")
        dt_inicio_ajustado = dt_inicio - datetime.timedelta(seconds=1)
        # lt é exclusivo, então usamos o início do dia seguinte
        dt_fim = datetime.datetime.strptime(data_fim, "%Y-%m-%d")
        dt_fim_ajustado = dt_fim + datetime.timedelta(days=1)

        conditions.append({
            "field": campo_filtro,
            "expression": "gt",
            "values": [dt_inicio_ajustado.strftime("%Y-%m-%dT%H:%M:%S") + "-03:00"],
        })
        conditions.append({
            "field": campo_filtro,
            "expression": "lt",
            "values": [dt_fim_ajustado.strftime("%Y-%m-%dT%H:%M:%S") + "-03:00"],
        })

    # Campos nativos que a API não suporta no search — filtrar localmente
    CAMPOS_LOCAIS = {"status", "origin", "source"}

    # Filtros adicionais por campos personalizados (múltiplos)
    filtros_extras = data.get("filtros_extras", [])
    for filtro in filtros_extras:
        campo_extra = (filtro.get("campo_extra") or "").strip()
        valor_extra = (filtro.get("valor_extra") or "").strip()
        tipo_busca = filtro.get("tipo_busca", "match")
        if campo_extra and valor_extra:
            if tipo_busca == "not_contains" or campo_extra in CAMPOS_LOCAIS:
                filtros_locais.append({
                    "field": campo_extra,
                    "value": valor_extra,
                    "tipo": tipo_busca,
                    "native": campo_extra in CAMPOS_LOCAIS,
                })
            else:
                conditions.append({
                    "field": campo_extra,
                    "expression": tipo_busca,
                    "values": [valor_extra],
                })

    # Compatibilidade: campo_extra único (legado)
    campo_extra = data.get("campo_extra", "").strip()
    valor_extra = data.get("valor_extra", "").strip()
    tipo_busca = data.get("tipo_busca", "match")
    if campo_extra and valor_extra:
        if tipo_busca == "not_contains" or campo_extra in CAMPOS_LOCAIS:
            filtros_locais.append({
                "field": campo_extra,
                "value": valor_extra,
                "tipo": tipo_busca,
                "native": campo_extra in CAMPOS_LOCAIS,
            })
        else:
            conditions.append({
                "field": campo_extra,
                "expression": tipo_busca,
                "values": [valor_extra],
            })

    return conditions, filtros_locais


def aplicar_filtros_locais(deals, filtros_locais, cf_map):
    """Aplica filtros locais nos deals: campos nativos (status, origin, source) e not_contains."""
    if not filtros_locais:
        return deals

    resultado = []
    for deal in deals:
        incluir = True
        for filtro in filtros_locais:
            campo = filtro["field"]
            valor = filtro["value"]
            tipo = filtro.get("tipo", "match")
            is_native = filtro.get("native", False)

            if is_native:
                # Campo nativo do deal (status, origin, source)
                deal_val = str(deal.get(campo) or "")
                if tipo == "not_contains":
                    if valor.lower() in deal_val.lower():
                        incluir = False
                else:
                    # match / contains / eq — comparação case-insensitive
                    if deal_val.upper() != valor.upper():
                        incluir = False
            else:
                # Custom field — filtro not_contains
                valor_lower = valor.lower()
                for ecf in deal.get("entityCustomFields") or []:
                    cf_id = ecf.get("id")
                    cf_nome = cf_map.get(cf_id, str(cf_id))
                    if cf_nome == campo or str(cf_id) == campo:
                        text_val = (ecf.get("textValue") or "").lower()
                        if valor_lower in text_val:
                            incluir = False
                            break

            if not incluir:
                break

        if incluir:
            resultado.append(deal)

    return resultado


@app.route("/buscar", methods=["POST"])
def buscar():
    """Busca deals e retorna JSON."""
    data = request.json
    api_key = data.get("api_key", "").strip()

    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    conditions, filtros_locais = montar_conditions(data)
    if not conditions and not filtros_locais:
        return jsonify({"error": "Informe pelo menos um filtro (datas ou campo personalizado)"}), 400

    try:
        # Se só há filtros locais (ex: status=LOST sem datas), precisamos de uma condição mínima para a API
        if not conditions and filtros_locais:
            conditions = [
                {"field": "dateCreated", "expression": "gt", "values": ["2000-01-01T00:00:00-03:00"]},
            ]

        deals = buscar_deals(api_key, conditions)
        cf_map = buscar_custom_fields(api_key)
        stage_map = buscar_stages(api_key)

        # Aplicar filtros locais (not_contains) antes de processar
        if filtros_locais:
            deals = aplicar_filtros_locais(deals, filtros_locais, cf_map)

        resumo = []
        for d in deals:
            # Custom fields do deal
            custom_fields = {}
            for ecf in d.get("entityCustomFields") or []:
                cf_id = ecf.get("id")
                cf_nome = cf_map.get(cf_id, cf_id)
                custom_fields[cf_nome] = ecf.get("textValue", "")

            stage_id = (d.get("stage") or {}).get("id", "")
            stage_name = stage_map.get(stage_id, str(stage_id)) if stage_id else ""

            resumo.append({
                "id": d.get("id"),
                "name": d.get("name", ""),
                "status": d.get("status", ""),
                "stage_id": stage_id,
                "stage_name": stage_name,
                "responsible_id": (d.get("responsible") or {}).get("id", ""),
                "contact_id": ((d.get("contacts") or [{}])[0]).get("id", "") if d.get("contacts") else "",
                "createdDate": d.get("dateCreated", ""),
                "origin": d.get("origin", ""),
                "source": d.get("source", ""),
                "customFields": custom_fields,
            })

        # Análise: contar leads por stage
        stage_counts = {}
        for d in resumo:
            sn = d["stage_name"] or "Sem etapa"
            stage_counts[sn] = stage_counts.get(sn, 0) + 1

        # Listar TODOS os stages disponíveis (não só os dos deals retornados)
        stage_names_list = sorted(set(stage_map.values()))

        return jsonify({
            "total": len(deals),
            "deals": resumo,
            "customFieldNames": list(cf_map.values()),
            "stageNames": stage_names_list,
            "stageCounts": stage_counts,
        })
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/deal_activities/<int:deal_id>", methods=["POST"])
def deal_activities(deal_id):
    """Retorna as atividades de um deal específico."""
    data = request.json or {}
    api_key = data.get("api_key", "").strip()
    if not api_key:
        return jsonify({"error": "API Key não informada"}), 400

    try:
        resp = requests.get(
            f"{API_BASE}/deals/{deal_id}/activities",
            headers=moskit_headers(api_key),
            timeout=15,
        )
        resp.raise_for_status()
        atividades = resp.json()
        return jsonify([
            {
                "type_id": (a.get("type") or {}).get("id"),
                "title": a.get("title", ""),
                "doneDate": a.get("doneDate"),
                "dateCreated": a.get("dateCreated"),
            }
            for a in atividades
        ])
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/exportar", methods=["POST"])
def exportar():
    """Busca deals e exporta para Excel."""
    data = request.json
    api_key = data.get("api_key", "").strip()
    data_inicio = data.get("data_inicio", "").strip()
    data_fim = data.get("data_fim", "").strip()

    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    conditions, filtros_locais = montar_conditions(data)
    if not conditions:
        return jsonify({"error": "Informe pelo menos um filtro"}), 400

    try:
        deals = buscar_deals(api_key, conditions)
        cf_map = buscar_custom_fields(api_key)

        # Aplicar filtros locais (not_contains)
        if filtros_locais:
            deals = aplicar_filtros_locais(deals, filtros_locais, cf_map)

        if not deals:
            return jsonify({"error": "Nenhum deal encontrado para os filtros informados"}), 404

        excel_file = gerar_excel(deals, cf_map)
        nome_arquivo = f"leads_moskit_{data_inicio}_a_{data_fim}.xlsx"

        return send_file(
            excel_file,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=nome_arquivo,
        )
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/motivos-perda", methods=["POST"])
def motivos_perda():
    """Retorna lista de motivos de perda distintos dos deals LOST da busca atual."""
    data = request.json
    api_key = data.get("api_key", "").strip()
    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    conditions, filtros_locais = montar_conditions(data)
    if not conditions and not filtros_locais:
        conditions = [
            {"field": "dateCreated", "expression": "gt", "values": ["2000-01-01T00:00:00-03:00"]},
        ]

    try:
        deals = buscar_deals(api_key, conditions)
        cf_map = buscar_custom_fields(api_key)

        if filtros_locais:
            deals = aplicar_filtros_locais(deals, filtros_locais, cf_map)

        motivos = set()
        for d in deals:
            if d.get("status") == "LOST":
                lr = d.get("lostReason")
                if isinstance(lr, dict):
                    nome = lr.get("name", "").strip()
                elif lr:
                    nome = str(lr).strip()
                else:
                    nome = ""
                motivos.add(nome if nome else "(sem motivo)")

        return jsonify(sorted(motivos))
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/exportar-desqualificados", methods=["POST"])
def exportar_desqualificados():
    """Exporta para Excel os leads com status LOST, com filtro opcional por motivo de perda."""
    data = request.json
    api_key = data.get("api_key", "").strip()
    data_inicio = data.get("data_inicio", "").strip()
    data_fim = data.get("data_fim", "").strip()
    motivos_filtro = data.get("motivos", [])  # lista de strings; vazia = todos

    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    conditions, filtros_locais = montar_conditions(data)
    if not conditions and not filtros_locais:
        conditions = [
            {"field": "dateCreated", "expression": "gt", "values": ["2000-01-01T00:00:00-03:00"]},
        ]

    try:
        deals = buscar_deals(api_key, conditions)
        cf_map = buscar_custom_fields(api_key)

        if filtros_locais:
            deals = aplicar_filtros_locais(deals, filtros_locais, cf_map)

        # Filtrar apenas LOST
        deals_lost = [d for d in deals if d.get("status") == "LOST"]

        # Filtrar por motivo se especificado
        if motivos_filtro:
            motivos_set = set(motivos_filtro)
            def get_motivo(d):
                lr = d.get("lostReason")
                if isinstance(lr, dict):
                    return lr.get("name", "").strip() or "(sem motivo)"
                return str(lr).strip() if lr else "(sem motivo)"

            deals_lost = [d for d in deals_lost if get_motivo(d) in motivos_set]

        if not deals_lost:
            return jsonify({"error": "Nenhum lead desqualificado encontrado para os filtros informados"}), 404

        excel_file = gerar_excel(deals_lost, cf_map, titulo="Desqualificados")
        sufixo_motivo = "_" + "_".join(m[:20] for m in motivos_filtro[:3]) if motivos_filtro else ""
        nome_arquivo = f"leads_desqualificados{sufixo_motivo}_{data_inicio}_a_{data_fim}.xlsx"

        return send_file(
            excel_file,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=nome_arquivo,
        )
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


ORDEM_FUNIL = [
    462709,  # Novo Lead
    457523,  # Sem Contato
    475012,  # Em Contato
    457521,  # Pré-Qualificação (Abertura)
    457522,  # Qualificação
    457659,  # Webinário
    475013,  # StandBy
    457658,  # Conferência Agendada (Negociação)
    457660,  # Conferência Realizada (Fechamento)
]

ETAPA_PRE_QUAL = 457521
ETAPA_QUAL = 457522


def passou_pela_etapa(deal, etapa_id):
    """Retorna True se o deal está na etapa_id ou em etapa posterior do funil."""
    stage_id = (deal.get("stage") or {}).get("id")
    if not stage_id:
        return False
    try:
        idx_etapa = ORDEM_FUNIL.index(etapa_id)
        idx_lead = ORDEM_FUNIL.index(int(stage_id))
        return idx_lead >= idx_etapa
    except ValueError:
        return False


def passou_apenas_pre_qual(deal):
    """Passou por Pré-Qualificação mas NÃO chegou à Qualificação."""
    return passou_pela_etapa(deal, ETAPA_PRE_QUAL) and not passou_pela_etapa(deal, ETAPA_QUAL)


@app.route("/exportar-pre-qualificados", methods=["POST"])
def exportar_pre_qualificados():
    """Exporta leads que passaram por Pré-Qualificação mas não chegaram à Qualificação."""
    data = request.json
    api_key = data.get("api_key", "").strip()
    data_inicio = data.get("data_inicio", "").strip()
    data_fim = data.get("data_fim", "").strip()

    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    conditions, filtros_locais = montar_conditions(data)
    if not conditions and not filtros_locais:
        conditions = [{"field": "dateCreated", "expression": "gt", "values": ["2000-01-01T00:00:00-03:00"]}]

    try:
        deals = buscar_deals(api_key, conditions)
        cf_map = buscar_custom_fields(api_key)
        if filtros_locais:
            deals = aplicar_filtros_locais(deals, filtros_locais, cf_map)

        filtrados = [d for d in deals if passou_apenas_pre_qual(d)]
        if not filtrados:
            return jsonify({"error": "Nenhum lead em Pré-Qualificação encontrado"}), 404

        excel_file = gerar_excel(filtrados, cf_map, titulo="Pré-Qualificação")
        return send_file(
            excel_file,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"leads_pre_qualificados_{data_inicio}_a_{data_fim}.xlsx",
        )
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/exportar-qualificados", methods=["POST"])
def exportar_qualificados():
    """Exporta leads que passaram pela etapa de Qualificação (457522) ou além."""
    data = request.json
    api_key = data.get("api_key", "").strip()
    data_inicio = data.get("data_inicio", "").strip()
    data_fim = data.get("data_fim", "").strip()

    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    conditions, filtros_locais = montar_conditions(data)
    if not conditions and not filtros_locais:
        conditions = [{"field": "dateCreated", "expression": "gt", "values": ["2000-01-01T00:00:00-03:00"]}]

    try:
        deals = buscar_deals(api_key, conditions)
        cf_map = buscar_custom_fields(api_key)
        if filtros_locais:
            deals = aplicar_filtros_locais(deals, filtros_locais, cf_map)

        filtrados = [d for d in deals if passou_pela_etapa(d, ETAPA_QUAL)]
        if not filtrados:
            return jsonify({"error": "Nenhum lead em Qualificação encontrado"}), 404

        excel_file = gerar_excel(filtrados, cf_map, titulo="Qualificação")
        return send_file(
            excel_file,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"leads_qualificados_{data_inicio}_a_{data_fim}.xlsx",
        )
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/debug-cf", methods=["POST"])
def debug_cf():
    """Retorna todos os custom fields disponíveis (id + nome) para diagnóstico."""
    data = request.json
    api_key = data.get("api_key", "").strip()
    if not api_key:
        return jsonify({"error": "API Key obrigatória"}), 400
    try:
        cf_map = buscar_custom_fields(api_key)
        return jsonify(sorted(cf_map.items(), key=lambda x: x[1]))
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/exportar-json", methods=["POST"])
def exportar_json():
    """Exporta deals em JSON no formato [{meta_leadgen_id, nome, moskit_id}].
    Aceita o parâmetro 'tipo': 'todos' | 'pre_qual' | 'qual' | 'desqualificados'.
    Para 'desqualificados' aceita também 'motivos' (lista de strings).
    """
    data = request.json
    api_key = data.get("api_key", "").strip()
    tipo = data.get("tipo", "todos")
    motivos_filtro = data.get("motivos", [])
    data_inicio = data.get("data_inicio", "").strip()
    data_fim = data.get("data_fim", "").strip()

    if not api_key:
        return jsonify({"error": "API Key é obrigatória"}), 400

    conditions, filtros_locais = montar_conditions(data)
    if not conditions and not filtros_locais:
        conditions = [{"field": "dateCreated", "expression": "gt", "values": ["2000-01-01T00:00:00-03:00"]}]

    try:
        deals = buscar_deals(api_key, conditions)
        cf_map = buscar_custom_fields(api_key)

        if filtros_locais:
            deals = aplicar_filtros_locais(deals, filtros_locais, cf_map)

        if tipo == "pre_qual":
            deals = [d for d in deals if passou_apenas_pre_qual(d)]
            nome_arquivo = f"leads_pre_qualificados_{data_inicio}_a_{data_fim}.json"
        elif tipo == "qual":
            deals = [d for d in deals if passou_pela_etapa(d, ETAPA_QUAL)]
            nome_arquivo = f"leads_qualificados_{data_inicio}_a_{data_fim}.json"
        elif tipo == "desqualificados":
            deals = [d for d in deals if d.get("status") == "LOST"]
            if motivos_filtro:
                motivos_set = set(motivos_filtro)
                def get_motivo(d):
                    lr = d.get("lostReason")
                    if isinstance(lr, dict):
                        return lr.get("name", "").strip() or "(sem motivo)"
                    return str(lr).strip() if lr else "(sem motivo)"
                deals = [d for d in deals if get_motivo(d) in motivos_set]
            sufixo = "_" + "_".join(m[:20] for m in motivos_filtro[:3]) if motivos_filtro else ""
            nome_arquivo = f"leads_desqualificados{sufixo}_{data_inicio}_a_{data_fim}.json"
        else:
            nome_arquivo = f"leads_moskit_{data_inicio}_a_{data_fim}.json"

        if not deals:
            return jsonify({"error": "Nenhum lead encontrado para os filtros informados"}), 404

        resultado = gerar_json_leads(deals, cf_map)
        json_str = "return " + json.dumps(resultado, ensure_ascii=False, indent=2) + ";"

        return Response(
            json_str,
            mimetype="application/json",
            headers={"Content-Disposition": f"attachment; filename={nome_arquivo}"},
        )
    except requests.exceptions.HTTPError as e:
        return jsonify({"error": f"Erro na API: {e.response.status_code} - {e.response.text}"}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/exportar-analise", methods=["POST"])
def exportar_analise():
    """Gera Excel com os dados da análise por campos de agrupamento (suporta múltiplas etapas de funil)."""
    data = request.json or {}
    campos_agrupamento = data.get("camposAgrupamento", [])
    if not campos_agrupamento:
        campo_antigo = data.get("campoAnuncio", "")
        if campo_antigo:
            campos_agrupamento = [campo_antigo]
    # Suporte a múltiplas etapas (novo formato) e etapa única (legado)
    etapas_marcadas = data.get("etapasMarcadas", [])
    if not etapas_marcadas:
        etapa_legado = data.get("etapaSelecionada", "")
        if etapa_legado:
            etapas_marcadas = [etapa_legado]
    total_geral = data.get("totalGeral", 0)
    linhas = data.get("linhas", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Análise de Qualificação"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="8E44AD", end_color="8E44AD", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    bold_font = Font(bold=True, size=11)

    # Resumo
    ws.cell(row=1, column=1, value="Resumo da Análise").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="Campos analisados:").font = bold_font
    ws.cell(row=2, column=2, value=" ▸ ".join(campos_agrupamento))
    ws.cell(row=3, column=1, value="Total de Leads:").font = bold_font
    ws.cell(row=3, column=2, value=total_geral)
    ws.cell(row=4, column=1, value="Etapas do funil:").font = bold_font
    ws.cell(row=4, column=2, value=" → ".join(etapas_marcadas) if etapas_marcadas else "—")

    # Tabela
    start_row = 6
    headers = list(campos_agrupamento) + ["Total Leads"]
    for etapa in etapas_marcadas:
        headers += [etapa, f"% → {etapa}"]
    headers += ["Ganhos", "% Ganho", "Perdidos", "% Perdido"]

    for col_idx, titulo in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=titulo)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for i, item in enumerate(linhas):
        dados = item[1] if isinstance(item, list) else item
        valores = dados.get("valores", []) if isinstance(dados, dict) else []
        data_row = start_row + 1 + i

        for col_idx, val in enumerate(valores, 1):
            ws.cell(row=data_row, column=col_idx, value=val).border = thin_border

        col_total = len(campos_agrupamento) + 1
        total = dados.get("total", 0) if isinstance(dados, dict) else 0
        ws.cell(row=data_row, column=col_total, value=total).border = thin_border

        por_etapa = dados.get("porEtapa", {}) if isinstance(dados, dict) else {}
        for j, etapa in enumerate(etapas_marcadas):
            n = por_etapa.get(etapa, dados.get("naEtapa", 0) if j == 0 else 0)
            pct = f"{((n / total) * 100):.1f}%" if total > 0 else "0.0%"
            col_e = col_total + 1 + j * 2
            ws.cell(row=data_row, column=col_e, value=n).border = thin_border
            ws.cell(row=data_row, column=col_e + 1, value=pct).border = thin_border

        col_desfecho = col_total + 1 + len(etapas_marcadas) * 2
        ganho   = dados.get("ganho", 0) if isinstance(dados, dict) else 0
        perdido = dados.get("perdido", 0) if isinstance(dados, dict) else 0
        pct_g = f"{((ganho   / total) * 100):.1f}%" if total > 0 else "0.0%"
        pct_p = f"{((perdido / total) * 100):.1f}%" if total > 0 else "0.0%"
        ws.cell(row=data_row, column=col_desfecho,     value=ganho).border   = thin_border
        ws.cell(row=data_row, column=col_desfecho + 1, value=pct_g).border   = thin_border
        ws.cell(row=data_row, column=col_desfecho + 2, value=perdido).border = thin_border
        ws.cell(row=data_row, column=col_desfecho + 3, value=pct_p).border   = thin_border

    # Ajustar largura
    for col_idx in range(1, len(headers) + 1):
        max_len = 0
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        for r in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws.max_row):
            for cell in r:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nome_arquivo = "analise_" + "_".join(campos_agrupamento)[:100] + ".xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nome_arquivo,
    )


if __name__ == "__main__":
    app.run(debug=False, port=5001)
